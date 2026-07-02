# -*- coding: utf-8 -*-
"""Builds the EduNexus Project Evaluation presentation (.pptx)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ----------------------------------------------------------------------------
# Theme
# ----------------------------------------------------------------------------
NAVY   = RGBColor(0x1F, 0x2A, 0x44)
NAVY2  = RGBColor(0x2B, 0x3A, 0x5C)
TEAL   = RGBColor(0x14, 0xB8, 0xA6)
TEAL_D = RGBColor(0x0E, 0x7C, 0x71)
AMBER  = RGBColor(0xF2, 0xA9, 0x00)
LIGHT  = RGBColor(0xF4, 0xF6, 0xF9)
GREY   = RGBColor(0x5A, 0x63, 0x72)
DARK   = RGBColor(0x22, 0x28, 0x33)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
BOXBLUE= RGBColor(0xE7, 0xEE, 0xF7)
BOXTEAL= RGBColor(0xD7, 0xF2, 0xEE)
BOXAMB = RGBColor(0xFD, 0xF0, 0xD6)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

_slide_no = 0

def new_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s

def _no_line(sp):
    sp.line.fill.background()

def set_text(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, font="Segoe UI"):
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run(); run.text = text
    f = run.font
    f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = font
    return p

def textbox(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, font="Segoe UI"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    set_text(tf, text, size, color, bold, align, font)
    return tb

def titlebar(slide, title, kicker=None):
    global _slide_no
    _slide_no += 1
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(1.15))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; _no_line(bar); bar.shadow.inherit=False
    acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.15), SW, Inches(0.07))
    acc.fill.solid(); acc.fill.fore_color.rgb = TEAL; _no_line(acc); acc.shadow.inherit=False
    tf = bar.text_frame; tf.word_wrap=True
    tf.margin_left=Inches(0.5); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    if kicker:
        p0=tf.paragraphs[0]; r0=p0.add_run(); r0.text=kicker
        r0.font.size=Pt(11); r0.font.bold=True; r0.font.color.rgb=TEAL; r0.font.name="Segoe UI"
        p1=tf.add_paragraph()
    else:
        p1=tf.paragraphs[0]
    r1=p1.add_run(); r1.text=title
    r1.font.size=Pt(26); r1.font.bold=True; r1.font.color.rgb=WHITE; r1.font.name="Segoe UI"
    # slide number
    num = slide.shapes.add_textbox(SW-Inches(1.0), Inches(0.32), Inches(0.8), Inches(0.5))
    set_text(num.text_frame, str(_slide_no), 12, TEAL, bold=True, align=PP_ALIGN.RIGHT)

def bullets(slide, items, x, y, w, h, size=16, gap=6, color=DARK, marker_color=TEAL):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap=True
    for i,(lead, body) in enumerate(items):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run(); r.text="▸ "
        r.font.size=Pt(size); r.font.bold=True; r.font.color.rgb=marker_color; r.font.name="Segoe UI"
        if lead:
            rl = p.add_run(); rl.text=lead+"  "
            rl.font.size=Pt(size); rl.font.bold=True; rl.font.color.rgb=NAVY; rl.font.name="Segoe UI"
        rb = p.add_run(); rb.text=body
        rb.font.size=Pt(size); rb.font.color.rgb=color; rb.font.name="Segoe UI"
    return tb

def box(slide, x, y, w, h, text, fill, fcolor=DARK, size=12, bold=False,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, line_w=1.0):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        _no_line(sp)
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit=False
    tf = sp.text_frame; tf.word_wrap=True
    tf.margin_left=Inches(0.05); tf.margin_right=Inches(0.05)
    tf.margin_top=Inches(0.03); tf.margin_bottom=Inches(0.03)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    set_text(tf, text, size, fcolor, bold, PP_ALIGN.CENTER)
    return sp

def connector(slide, x1, y1, x2, y2, color=GREY, width=1.5, arrow=True, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color; c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if arrow:
        tail = ln.makeelement(qn('a:tailEnd'), {'type':'triangle','w':'med','len':'med'})
        ln.append(tail)
    if dash:
        d = ln.makeelement(qn('a:prstDash'), {'val':'dash'})
        ln.append(d)
    c.shadow.inherit=False
    return c

def edge_label(slide, x, y, text, color=GREY, size=9, w=Inches(1.6)):
    tb = slide.shapes.add_textbox(x, y, w, Inches(0.3))
    tf=tb.text_frame; tf.word_wrap=True
    set_text(tf, text, size, color, align=PP_ALIGN.CENTER)
    return tb

def chip(slide, x, y, w, text, fill=TEAL, fcolor=WHITE, size=11):
    return box(slide, x, y, w, Inches(0.36), text, fill, fcolor, size, bold=True,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE)

# ============================================================================
# 1. TITLE
# ============================================================================
s = new_slide(NAVY)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.5), SW, Inches(0.06))
band.fill.solid(); band.fill.fore_color.rgb=TEAL; _no_line(band); band.shadow.inherit=False
textbox(s, Inches(1), Inches(0.7), Inches(11.3), Inches(0.5),
        "FINAL YEAR PROJECT  •  PROJECT EVALUATION", 14, TEAL, bold=True, align=PP_ALIGN.CENTER)
textbox(s, Inches(0.8), Inches(1.4), Inches(11.7), Inches(1.2),
        "EduNexus", 60, WHITE, bold=True, align=PP_ALIGN.CENTER)
textbox(s, Inches(1), Inches(2.7), Inches(11.3), Inches(1.0),
        "An AI-Powered School Management System", 26, RGBColor(0xCF,0xE9,0xE5),
        bold=True, align=PP_ALIGN.CENTER)
textbox(s, Inches(1), Inches(3.7), Inches(11.3), Inches(0.8),
        "Automating academic & administrative operations with intelligent, "
        "natural-language interaction", 15, RGBColor(0xB8,0xC2,0xD4), align=PP_ALIGN.CENTER)
# tech chips
techs = ["Next.js", "Node / Express", "MongoDB", "n8n AI Agent", "Redis + Bull", "JWT / RBAC"]
cx = Inches(1.85); cw=Inches(1.55)
for t in techs:
    chip(s, cx, Inches(5.0), cw, t, fill=NAVY2, fcolor=TEAL); cx += cw+Inches(0.12)
textbox(s, Inches(1), Inches(6.4), Inches(11.3), Inches(0.5),
        "Methodology: Iterative Waterfall", 13, RGBColor(0x9A,0xA6,0xBC), align=PP_ALIGN.CENTER)

# ============================================================================
# 2. PROBLEM STATEMENT
# ============================================================================
s = new_slide(LIGHT)
titlebar(s, "Problem Statement", "THE GAP")
bullets(s, [
    ("", "Educational institutions still rely on outdated or disconnected software systems that rarely talk to each other."),
    ("", "Academic and administrative processes (attendance, marks, fees, timetables, complaints) demand heavy manual effort — they are labor-intensive, error-prone and time-consuming."),
    ("", "Data is locked in silos, so management cannot get quick, reliable answers or insights for decision-making."),
    ("", "Non-technical staff struggle with complex dashboards and reports."),
], Inches(0.7), Inches(1.55), Inches(7.3), Inches(4.0), size=17, gap=12)
# solution callout
box(s, Inches(8.4), Inches(1.7), Inches(4.3), Inches(4.0),
    "OUR ANSWER\n\nA single, integrated, AI-powered platform that automates "
    "academic & administrative operations — and lets any user simply ASK questions "
    "in natural language to get insights, summaries, charts and downloadable reports.",
    fill=NAVY, fcolor=WHITE, size=16, bold=False)
box(s, Inches(8.4), Inches(1.7), Inches(4.3), Inches(0.55), "OUR ANSWER", TEAL, WHITE, 15, bold=True)

# ============================================================================
# 3. RELATED PROJECTS
# ============================================================================
s = new_slide(LIGHT)
titlebar(s, "Related Projects / Existing Systems", "MARKET STUDY")
projects = [
    ("SIMSIN ERP", "Comprehensive school ERP; modular but heavy to configure and largely menu-driven with no AI assistance."),
    ("Edusuite", "Cloud academic suite covering admissions, exams & fees; limited conversational analytics."),
    ("Fedena", "Popular open-source/SaaS school ERP; broad modules but reporting is static and report-building is manual."),
    ("Entab", "Established Indian school management & parent-app vendor; strong on operations, weak on natural-language insight."),
    ("Skoolie", "Lightweight school app focused on communication & attendance; narrow administrative coverage."),
]
y = Inches(1.5)
for i,(name, desc) in enumerate(projects):
    box(s, Inches(0.7), y, Inches(2.6), Inches(0.92), name, TEAL_D, WHITE, 15, bold=True)
    box(s, Inches(3.45), y, Inches(9.2), Inches(0.92), desc, WHITE, DARK, 13.5,
        line=RGBColor(0xD5,0xDB,0xE5))
    y += Inches(1.05)
textbox(s, Inches(0.7), Inches(6.95), Inches(12), Inches(0.4),
        "Gap addressed by EduNexus: an end-to-end, role-aware platform with a built-in AI query engine — no manual report building.",
        12.5, GREY, bold=True)

# ============================================================================
# 4. OUR PROJECT
# ============================================================================
s = new_slide(LIGHT)
titlebar(s, "Our Project — EduNexus", "OVERVIEW")
textbox(s, Inches(0.7), Inches(1.45), Inches(12), Inches(1.2),
        "EduNexus is a unified, cloud-ready School Management System that replaces fragmented manual workflows "
        "with one connected platform. It combines a role-based operational core (attendance, examinations, fees, "
        "timetables, homework, complaints, transport & communication) with an AI query engine that turns plain "
        "natural-language questions into secure, role-scoped insights, charts and downloadable reports.", 15, DARK)
cards = [
    ("Unified Core", "One system for academics, finance, HR & operations — no more disconnected tools.", BOXBLUE, NAVY),
    ("AI Query Engine", "Ask in natural language; an n8n-driven agent answers within each user's permissions.", BOXTEAL, TEAL_D),
    ("Role-Based Access", "8 roles — Admin, Principal, Finance, HR, Reception, Teacher, Student, Parent.", BOXAMB, RGBColor(0x9A,0x6B,0x00)),
    ("Real-Time Insight", "Live dashboards, push notifications & secure in-house messaging.", BOXBLUE, NAVY),
]
cx = Inches(0.7); cw=Inches(2.95)
for title_c, body_c, fill_c, fc in cards:
    box(s, cx, Inches(3.05), cw, Inches(0.55), title_c, fc, WHITE, 14, bold=True)
    box(s, cx, Inches(3.62), cw, Inches(1.85), body_c, fill_c, DARK, 13)
    cx += cw+Inches(0.18)
# stack strip
textbox(s, Inches(0.7), Inches(5.75), Inches(12), Inches(0.4), "Technology Stack", 14, NAVY, bold=True)
stack = ["Frontend: Next.js / React", "Backend: Node.js + Express", "DB: MongoDB (Mongoose, ~60 models)",
         "AI Orchestration: n8n Agent", "Queue/Cache: Redis + Bull", "Auth: JWT + RBAC"]
cx=Inches(0.7); cw=Inches(2.0)
for i,t in enumerate(stack):
    box(s, cx, Inches(6.2), cw, Inches(0.8), t, WHITE, DARK, 11.5, line=TEAL, line_w=1.25)
    cx += cw+Inches(0.05)

# ============================================================================
# 5. METHODOLOGY
# ============================================================================
s = new_slide(LIGHT)
titlebar(s, "Methodology — Iterative Waterfall", "PROCESS MODEL")
textbox(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.7),
        "We followed the Iterative Waterfall model: classic Waterfall phases executed in sequence, but with "
        "feedback loops that let us return to an earlier phase to refine requirements or fix defects each iteration.",
        14.5, DARK)
phases = ["Requirements", "Design", "Implementation", "Testing", "Deployment", "Maintenance"]
n=len(phases); pw=Inches(1.85); gap=Inches(0.12)
total = n*pw + (n-1)*gap
x = Emu(int((SW - total)/2)); y=Inches(2.7)
centers=[]
for i,ph in enumerate(phases):
    b=box(s, x, y, pw, Inches(0.95), ph, TEAL_D if i%2==0 else NAVY, WHITE, 14, bold=True)
    centers.append((x+Emu(int(pw/2)), y))
    if i < n-1:
        connector(s, x+pw, y+Emu(int(Inches(0.95)/2)), x+pw+gap, y+Emu(int(Inches(0.95)/2)), color=AMBER, width=2.5)
    x = Emu(int(x + pw + gap))
# feedback loop arcs (dashed back-arrows)
for i in range(n-1, 0, -1):
    if i in (1,2,3,4):
        cx_from = centers[i][0]; cx_to = centers[i-1][0]
        connector(s, cx_from, y+Inches(0.95), cx_to, y+Inches(0.95)+Inches(0.5),
                  color=GREY, width=1.25, dash=True)
textbox(s, Inches(0.7), Inches(4.55), Inches(12), Inches(0.4),
        "Dashed arrows = feedback loops: each test/review cycle can revisit earlier phases (iterative refinement).",
        12, GREY, bold=True)
why = [
    ("Why this model:", "Requirements were largely known up-front (school operations are well understood), so a structured Waterfall fit."),
    ("Iteration added:", "AI features, new roles and modules were refined across cycles based on supervisor & user feedback."),
    ("Risk control:", "Each phase is verified before the next, reducing rework while still allowing controlled change."),
]
bullets(s, why, Inches(0.7), Inches(5.15), Inches(12), Inches(2.0), size=14.5, gap=9)

# ============================================================================
# 6. OBJECTIVE
# ============================================================================
s = new_slide(LIGHT)
titlebar(s, "Objective", "GOAL")
box(s, Inches(1.2), Inches(1.7), Inches(10.9), Inches(1.7),
    "To design and develop an AI-powered, role-based School Management System that automates core "
    "academic and administrative operations and enables every user to obtain accurate, secure, "
    "permission-aware insights and reports through simple natural-language interaction.",
    NAVY, WHITE, 19, bold=True)
textbox(s, Inches(0.7), Inches(3.8), Inches(12), Inches(0.45), "Supporting objectives", 15, NAVY, bold=True)
sub = [
    ("", "Eliminate manual, paper-based and disconnected processes through one integrated platform."),
    ("", "Enforce strict role-based access so users see only what they are permitted to."),
    ("", "Provide an AI query engine that returns insights, summaries, charts and downloadable reports on demand."),
    ("", "Improve transparency and communication between school, teachers, students and parents."),
    ("", "Deliver a responsive, reliable and maintainable system suitable for real institutions."),
]
bullets(s, sub, Inches(0.9), Inches(4.25), Inches(11.5), Inches(2.8), size=15, gap=9)

# ============================================================================
# 7. PROJECT SCOPE (Pakistan)
# ============================================================================
s = new_slide(LIGHT)
titlebar(s, "Project Scope (Pakistan)", "SCOPE")
textbox(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.95),
        "EduNexus targets schools and academic institutions across Pakistan — from single-campus private "
        "schools to multi-section setups — that need an affordable, locally-relevant alternative to fragmented "
        "registers and costly foreign ERPs.", 15, DARK)
sc = [
    ("Target users:", "Private schools, academies and college sections in Pakistan; admin, faculty, students & parents."),
    ("Local fit:", "Supports Pakistani academic patterns — term/result cards, fee + concession structures, transport fees, and parent communication."),
    ("Affordability:", "Cloud-ready and open-source-stack based to keep deployment cost low for local institutions."),
    ("Accessibility:", "Responsive web on desktop, tablet & mobile — usable on the devices Pakistani parents & staff already own."),
    ("Coverage:", "Attendance, exams/report cards, fees & transport, timetables, homework, complaints, notifications, messaging and AI insights."),
    ("Out of scope:", "Hardware (biometric/RFID) integration, native mobile apps and offline mode are future enhancements."),
]
bullets(s, sc, Inches(0.7), Inches(2.5), Inches(12), Inches(4.5), size=15, gap=11)

# ============================================================================
# 8. FUNCTIONAL REQUIREMENTS (2 slides)
# ============================================================================
fr_all = [
    ("FR-1 — Authentication & RBAC",
     "Secure JWT-based authentication with encrypted password validation; dashboards and access load per the role and permissions assigned to each user."),
    ("FR-2 — User & Role Management",
     "Admin can create, update, delete and assign roles: Admin, Principal, Finance, HR, Reception, Teacher, Student, Parent."),
    ("FR-3 — Attendance Management",
     "Teachers record and edit daily attendance; Admin, Principal, HR, Parents, Students & Reception view summaries or details per access level."),
    ("FR-4 — Examination & Report Cards",
     "Teachers enter marks, grades & comments; report cards auto-generate; Principals approve final batches; Students/Parents view or download published cards."),
    ("FR-5 — Timetable Management",
     "Admin uploads/edits class timetables; Teachers, Students, Parents, Principal & Reception access updated timetables with role-appropriate visibility."),
    ("FR-6 — Complaints & Feedback",
     "Teachers, Parents & Students submit complaints/feedback; Principal tracks, assigns & resolves; complainant gets real-time status updates."),
    ("FR-7 — Homework Management",
     "Teachers post homework & change status; Students submit digitally; Parents view homework & status; Principal audits posting & workload patterns."),
    ("FR-8 — Fee & Finance Management",
     "Principal/Finance manage fee structure, payments, concessions, receipts, financial reports & transport payments; Parents/Students/Reception view or verify status per permissions."),
    ("FR-9 — AI Query Engine",
     "AI-powered engine takes natural-language questions and returns insights, summaries, charts or downloadable reports based on the user's role permissions."),
    ("FR-10 — Notifications & Messaging",
     "Principal pushes notifications (attendance, grades, announcements, fee reminders, holidays, staff alerts) plus secure, monitored in-house messaging by permission."),
    ("FR-11 — Transport Management",
     "Finance manages transport payment records; Reception checks enrollment & payment status; Parents/Students view transport details; Admin configures routes."),
]
def fr_slide(title, items, start):
    s = new_slide(LIGHT)
    titlebar(s, title, "FUNCTIONAL REQUIREMENTS")
    y=Inches(1.45)
    for name, desc in items:
        box(s, Inches(0.6), y, Inches(3.4), Inches(0.85), name, NAVY, WHITE, 13, bold=True)
        box(s, Inches(4.1), y, Inches(8.6), Inches(0.85), desc, WHITE, DARK, 12.5,
            line=RGBColor(0xD5,0xDB,0xE5))
        y += Inches(0.95)
fr_slide("Functional Requirements (1 of 2)", fr_all[:6], 0)
fr_slide("Functional Requirements (2 of 2)", fr_all[6:], 6)

# ============================================================================
# 9. NON-FUNCTIONAL REQUIREMENTS
# ============================================================================
nfr = [
    ("NFR-1 — Performance",
     "Key metrics refresh at least every 5 minutes; supports ≥100 concurrent users while keeping dashboard, report, query and AI-insight response times under 90 seconds."),
    ("NFR-2 — Security",
     "Encrypts communications over HTTPS, enforces JWT authentication and strictly applies role-based access control for all user categories."),
    ("NFR-3 — Compatibility / Responsiveness",
     "Responsive interface that works across desktop, tablet and mobile to ensure accessibility and usability for all users."),
    ("NFR-4 — Usability",
     "Provides tool-tips, guided help, contextual prompts and personalized role-based dashboards for each role's frequent tasks."),
    ("NFR-5 — Reliability / Availability",
     "Maintains ≥90% uptime under normal load and ensures operational reliability of all core institutional processes."),
    ("NFR-6 — Recoverability",
     "Automatic backups with full restoration of all critical data, ensuring data integrity and recoverability."),
    ("NFR-7 — Maintainability",
     "Well-documented, modular codebase with extensive API documentation plus error-tracking, logging and monitoring tools."),
]
s = new_slide(LIGHT)
titlebar(s, "Non-Functional Requirements", "QUALITY ATTRIBUTES")
y=Inches(1.32)
for name, desc in nfr:
    box(s, Inches(0.6), y, Inches(3.2), Inches(0.72), name, TEAL_D, WHITE, 12.5, bold=True)
    box(s, Inches(3.9), y, Inches(8.8), Inches(0.72), desc, WHITE, DARK, 12,
        line=RGBColor(0xD5,0xDB,0xE5))
    y += Inches(0.80)

# ============================================================================
# CHAT MODULE INTRO
# ============================================================================
s = new_slide(NAVY)
textbox(s, Inches(1), Inches(2.6), Inches(11.3), Inches(0.6),
        "MODULE DEEP-DIVE", 16, TEAL, bold=True, align=PP_ALIGN.CENTER)
textbox(s, Inches(1), Inches(3.2), Inches(11.3), Inches(1.0),
        "AI Chat Module", 44, WHITE, bold=True, align=PP_ALIGN.CENTER)
textbox(s, Inches(1.5), Inches(4.4), Inches(10.3), Inches(1.0),
        "Natural-language query engine — the user asks, an n8n AI agent answers within "
        "their role permissions. Use-Case • Activity • DFD • Sequence",
        16, RGBColor(0xB8,0xC2,0xD4), align=PP_ALIGN.CENTER)

# ----------------------------------------------------------------------------
# Stick-figure actor
# ----------------------------------------------------------------------------
def actor(slide, cx, top, label, color=NAVY):
    head_d = Inches(0.36)
    h = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx-Emu(int(head_d/2)), top, head_d, head_d)
    h.fill.solid(); h.fill.fore_color.rgb=color; _no_line(h); h.shadow.inherit=False
    body_top = top+head_d
    connector(slide, cx, body_top, cx, body_top+Inches(0.55), color=color, width=2.5, arrow=False)
    connector(slide, cx-Inches(0.32), body_top+Inches(0.15), cx+Inches(0.32), body_top+Inches(0.15), color=color, width=2.5, arrow=False)
    connector(slide, cx, body_top+Inches(0.55), cx-Inches(0.28), body_top+Inches(0.95), color=color, width=2.5, arrow=False)
    connector(slide, cx, body_top+Inches(0.55), cx+Inches(0.28), body_top+Inches(0.95), color=color, width=2.5, arrow=False)
    lb = slide.shapes.add_textbox(cx-Inches(1.0), body_top+Inches(1.0), Inches(2.0), Inches(0.5))
    set_text(lb.text_frame, label, 12, color, bold=True, align=PP_ALIGN.CENTER)

# ============================================================================
# 10. USE CASE DIAGRAM (Chat Module)
# ============================================================================
s = new_slide(WHITE)
titlebar(s, "Use Case Diagram — Chat Module", "FR-9  AI QUERY ENGINE")
# system boundary
bnd = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.5), Inches(1.45), Inches(6.2), Inches(5.05))
bnd.fill.background(); bnd.line.color.rgb=NAVY; bnd.line.width=Pt(1.75); bnd.shadow.inherit=False
textbox(s, Inches(3.5), Inches(1.5), Inches(6.2), Inches(0.35), "AI Chat Module", 14, NAVY, bold=True, align=PP_ALIGN.CENTER)
# actors
actor(s, Inches(1.5), Inches(2.7), "User\n(any role)", NAVY)
actor(s, Inches(11.7), Inches(2.7), "n8n AI\nAgent", TEAL_D)
# use cases (single centered column => no line crosses a box)
OX=Inches(5.2); OW=Inches(2.85); OH=Inches(0.55)
uc_rows = [
    ("Send Natural-Language Query", Inches(1.80)),
    ("Receive Role-Scoped Insight / Chart", Inches(2.47)),
    ("View Chat History", Inches(3.14)),
    ("Resume / Open Session", Inches(3.81)),
    ("Download Report / Attachment", Inches(4.48)),
    ("Authenticate (JWT)", Inches(5.15)),
    ("Enforce Role-Based Access", Inches(5.82)),
]
ucy={}
for txt,y in uc_rows:
    box(s, OX, y, OW, OH, txt, BOXTEAL, TEAL_D, 11.5, bold=True, shape=MSO_SHAPE.OVAL)
    ucy[txt]=y
def umid(txt): return ucy[txt]+Emu(int(OH/2))
# user associations (approach left edge from far left -> never enters a box)
for txt in ["Send Natural-Language Query","View Chat History","Resume / Open Session","Download Report / Attachment"]:
    connector(s, Inches(1.7), Inches(3.4), OX, umid(txt), color=NAVY, width=1.25, arrow=False)
# agent association (approach right edge from far right)
connector(s, Inches(11.7), Inches(3.4), OX+OW, umid("Receive Role-Scoped Insight / Chart"),
          color=TEAL_D, width=1.25, arrow=False)
# «include» elbows routed in the empty right margin (no box crossings)
RX=OX+OW
def include_elbow(target, lane):
    sy=umid("Send Natural-Language Query"); ty=umid(target)
    connector(s, RX, sy, lane, sy, color=AMBER, width=1.1, arrow=False, dash=True)
    connector(s, lane, sy, lane, ty, color=AMBER, width=1.1, arrow=False, dash=True)
    connector(s, lane, ty, RX, ty, color=AMBER, width=1.1, arrow=True, dash=True)
include_elbow("Authenticate (JWT)", Inches(8.85))
include_elbow("Enforce Role-Based Access", Inches(9.20))
edge_label(s, Inches(8.35), Inches(3.55), "«include»", AMBER, 9.5, Inches(1.1))
textbox(s, Inches(0.5), Inches(6.65), Inches(12.4), Inches(0.7),
        "Description: Any authenticated User initiates queries and consumes role-scoped insights, history, sessions and "
        "downloads; the n8n AI Agent produces the answers. Every query «includes» JWT Authentication and Role-Based "
        "Access enforcement (dashed arrows).", 11, GREY)

# ============================================================================
# 11. ACTIVITY DIAGRAM (Chat Module)
# ============================================================================
s = new_slide(WHITE)
titlebar(s, "Activity Diagram — Chat Module", "QUERY LIFECYCLE")
def pill(x,y,w,h,t,fill,fc=WHITE,size=11,bold=True):
    return box(s,x,y,w,h,t,fill,fc,size,bold=bold,shape=MSO_SHAPE.ROUNDED_RECTANGLE)
def proc(x,y,w,h,t,fill=BOXBLUE,fc=DARK,size=10.5):
    return box(s,x,y,w,h,t,fill,fc,size,bold=False,shape=MSO_SHAPE.ROUNDED_RECTANGLE,line=NAVY,line_w=0.75)
def dec(x,y,w,h,t,size=10):
    return box(s,x,y,w,h,t,BOXAMB,RGBColor(0x8A,0x5D,0x00),size,bold=True,shape=MSO_SHAPE.DIAMOND,line=AMBER,line_w=1.0)

C1=Inches(0.8); C2=Inches(5.0); C3=Inches(9.4)
w=Inches(2.7); h=Inches(0.6)
cx1=C1+Emu(int(w/2)); cx2=C2+Emu(int(w/2)); cx3=C3+Emu(int(w/2))
def v(x,y1,y2,color=NAVY,arrow=True): connector(s,x,y1,x,y2,color=color,width=1.4,arrow=arrow)
def hl(x1,x2,y,color=NAVY,arrow=True): connector(s,x1,y,x2,y,color=color,width=1.4,arrow=arrow)

# ----- Column 1 -----
pill(cx1-Inches(0.6), Inches(1.35), Inches(1.2), Inches(0.45), "Start", TEAL, WHITE, 12)
proc(C1, Inches(1.95), w, h, "User types message (frontend)")
dec(cx1-Inches(0.8), Inches(2.75), Inches(1.6), Inches(0.9), "Valid?\n(empty / length)", size=9.5)
proc(C1, Inches(3.85), w, h, "Authenticate JWT + rate-limit")
proc(C1, Inches(4.55), w, h, "Ensure ChatSession (role scope)")
proc(C1, Inches(5.25), w, h, "Persist user message")
v(cx1, Inches(1.80), Inches(1.95))
v(cx1, Inches(1.95)+h, Inches(2.75))
v(cx1, Inches(2.75)+Inches(0.9), Inches(3.85)); edge_label(s, cx1+Inches(0.05), Inches(3.60), "Yes", TEAL_D, 9, Inches(0.5))
v(cx1, Inches(3.85)+h, Inches(4.55))
v(cx1, Inches(4.55)+h, Inches(5.25))
# invalid -> error terminal (clear of every box and of the C1->C2 hand-off lane)
EBX=Inches(3.6)
box(s, EBX, Inches(2.85), Inches(1.0), Inches(0.7), "Return 400\n& End", BOXAMB, RGBColor(0x8A,0x5D,0x00),
    9.5, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=AMBER, line_w=1.0)
hl(cx1+Inches(0.8), EBX, Inches(2.75)+Inches(0.45), color=AMBER)
edge_label(s, cx1+Inches(0.82), Inches(2.92), "No", AMBER, 9, Inches(0.6))
# hand-off: Persist user msg -> Cache decision (elbow in channel 1)
LANE1=Inches(4.75); PY=Inches(5.25)+Emu(int(h/2)); CDY=Inches(1.45)+Inches(0.475)
hl(C1+w, LANE1, PY, color=NAVY, arrow=False)
v(LANE1, PY, CDY, color=NAVY, arrow=False)
hl(LANE1, Inches(5.55), CDY, color=NAVY, arrow=True)
# ----- Column 2 -----
dec(cx2-Inches(0.8), Inches(1.45), Inches(1.6), Inches(0.95), "Cache hit?\n(Redis / mem)", size=9.5)
proc(C2, Inches(2.6), w, h, "Enqueue job (Bull) → worker")
proc(C2, Inches(3.3), w, h, "POST payload to n8n AI Agent")
proc(C2, Inches(4.0), w, h, "Agent: role scope + query DB")
proc(C2, Inches(4.7), w, h, "Agent returns reply / data / files")
v(cx2, Inches(1.45)+Inches(0.95), Inches(2.6)); edge_label(s, cx2+Inches(0.05), Inches(2.38), "No", AMBER, 9, Inches(0.5))
v(cx2, Inches(2.6)+h, Inches(3.3))
v(cx2, Inches(3.3)+h, Inches(4.0))
v(cx2, Inches(4.0)+h, Inches(4.7))
# ----- Column 3 -----
proc(C3, Inches(2.6), w, h, "Persist agent message + files")
proc(C3, Inches(3.3), w, h, "Update session (preview, time)")
proc(C3, Inches(4.0), w, h, "Cache response (TTL)")
proc(C3, Inches(4.7), w, h, "Return answer to user")
pill(cx3-Inches(0.6), Inches(5.5), Inches(1.2), Inches(0.45), "End", NAVY, WHITE, 12)
v(cx3, Inches(2.6)+h, Inches(3.3))
v(cx3, Inches(3.3)+h, Inches(4.0))
v(cx3, Inches(4.0)+h, Inches(4.7))
v(cx3, Inches(4.7)+h, Inches(5.5))
# hand-off: Agent returns -> Persist agent msg (elbow in channel 2, normal path)
LANE2=Inches(8.9); AY=Inches(4.7)+Emu(int(h/2)); B9=Inches(2.6)+Inches(0.40)
hl(C2+w, LANE2, AY, color=NAVY, arrow=False)
v(LANE2, AY, B9, color=NAVY, arrow=False)
hl(LANE2, C3, B9, color=NAVY, arrow=True)
# cache-hit shortcut: Cache decision (Yes) -> Persist agent msg (elbow, teal)
LANE3=Inches(8.5); CRY=Inches(1.45)+Inches(0.475); B9T=Inches(2.6)+Inches(0.18)
hl(cx2+Inches(0.8), LANE3, CRY, color=TEAL_D, arrow=False)
v(LANE3, CRY, B9T, color=TEAL_D, arrow=False)
hl(LANE3, C3, B9T, color=TEAL_D, arrow=True)
edge_label(s, Inches(7.55), Inches(1.18), "Yes (cached)", TEAL_D, 9, Inches(1.5))

# ============================================================================
# 12. DFD (Chat Module)
# ============================================================================
s = new_slide(WHITE)
titlebar(s, "Data Flow Diagram — Chat Module (Level 1)", "DFD")
# external entity
ee=box(s, Inches(0.6), Inches(3.0), Inches(2.0), Inches(1.1), "User\n(External Entity)", NAVY, WHITE, 12.5, bold=True, shape=MSO_SHAPE.RECTANGLE)
# process 1 Chat API
p1=box(s, Inches(3.6), Inches(2.95), Inches(2.4), Inches(1.2), "1.0\nChat API\n(Express)", BOXTEAL, TEAL_D, 13, bold=True, shape=MSO_SHAPE.OVAL)
# process 2 AI agent
p2=box(s, Inches(7.3), Inches(2.95), Inches(2.4), Inches(1.2), "2.0\nn8n AI Agent\n(query + RBAC)", BOXTEAL, TEAL_D, 12.5, bold=True, shape=MSO_SHAPE.OVAL)
# data stores (open rectangles)
def store(x,y,t):
    return box(s,x,y,Inches(2.7),Inches(0.55),t,BOXBLUE,NAVY,11,bold=True,line=NAVY,line_w=1.0,shape=MSO_SHAPE.RECTANGLE)
store(Inches(3.45), Inches(1.2), "D1  ChatSession / ChatMessage")
store(Inches(3.45), Inches(5.0), "D2  ChatFile (attachments)")
store(Inches(3.45), Inches(5.7), "D3  Cache (Redis / memory)")
store(Inches(10.1), Inches(1.5), "D4  Institution DB\n(~60 collections)")
# flows
connector(s, Inches(2.6), Inches(3.4), Inches(3.6), Inches(3.4), color=NAVY); edge_label(s, Inches(2.55), Inches(3.05), "NL query", NAVY, 9)
connector(s, Inches(3.6), Inches(3.75), Inches(2.6), Inches(3.75), color=TEAL_D); edge_label(s, Inches(2.55), Inches(3.78), "insight/report", TEAL_D, 9)
connector(s, Inches(6.0), Inches(3.4), Inches(7.3), Inches(3.4), color=NAVY); edge_label(s, Inches(6.05), Inches(3.05), "payload + role", NAVY, 9, Inches(1.4))
connector(s, Inches(7.3), Inches(3.75), Inches(6.0), Inches(3.75), color=TEAL_D); edge_label(s, Inches(6.05), Inches(3.78), "reply/data", TEAL_D, 9, Inches(1.3))
# p1 to stores
connector(s, Inches(4.8), Inches(2.95), Inches(4.8), Inches(1.75), color=GREY); edge_label(s, Inches(4.9), Inches(2.2), "persist msgs", GREY, 9, Inches(1.2))
connector(s, Inches(4.8), Inches(4.15), Inches(4.8), Inches(5.0), color=GREY); edge_label(s, Inches(4.9), Inches(4.5), "store files", GREY, 9, Inches(1.2))
connector(s, Inches(5.4), Inches(4.15), Inches(5.4), Inches(5.7), color=GREY, dash=True); edge_label(s, Inches(5.5), Inches(4.95), "get/set cache", GREY, 9, Inches(1.3))
# p2 to DB
connector(s, Inches(9.7), Inches(3.1), Inches(11.4), Inches(2.05), color=GREY); edge_label(s, Inches(9.9), Inches(2.4), "role-scoped read", GREY, 9, Inches(1.6))
textbox(s, Inches(0.6), Inches(6.7), Inches(12), Inches(0.6),
        "The User submits a natural-language query to the Chat API, which authenticates, persists the message and (on a cache miss) "
        "forwards a role-scoped payload to the n8n AI Agent. The agent reads the Institution DB within permissions and returns insight; "
        "responses are persisted, cached and delivered back to the user.", 10.5, GREY)

# ============================================================================
# 13. SEQUENCE DIAGRAM (Chat Module)
# ============================================================================
s = new_slide(WHITE)
titlebar(s, "Sequence Diagram — Chat Module", "RUNTIME FLOW")
lifelines = [
    ("User /\nFrontend", Inches(1.15)),
    ("Chat API\n(Express)", Inches(3.25)),
    ("Auth +\nRate Limit", Inches(5.15)),
    ("Cache\n(Redis)", Inches(6.95)),
    ("Queue\n(Bull)", Inches(8.55)),
    ("n8n AI\nAgent", Inches(10.3)),
    ("MongoDB", Inches(12.2)),
]
top=Inches(1.4); bottom=Inches(6.7)
xmap={}
for name,x in lifelines:
    box(s, x-Inches(0.85), top, Inches(1.7), Inches(0.6), name, NAVY, WHITE, 11, bold=True)
    connector(s, x, top+Inches(0.6), x, bottom, color=RGBColor(0xB8,0xC0,0xCC), width=1.0, arrow=False, dash=True)
    xmap[name]=x
def msg(y, a, b, text, color=NAVY, ret=False):
    x1=xmap[a]; x2=xmap[b]
    connector(s, x1, y, x2, y, color=color, width=1.5, arrow=True, dash=ret)
    mid=min(x1,x2)
    tb=s.shapes.add_textbox(mid+Inches(0.05), y-Inches(0.28), abs(int(x2-x1))+Inches(0.1) if x2>x1 else Inches(2.2), Inches(0.27))
    set_text(tb.text_frame, text, 9.5, color if not ret else GREY, align=PP_ALIGN.CENTER)
U,A,Au,C,Q,N,M = [l[0] for l in lifelines]
y=Inches(2.25); step=Inches(0.45)
msg(y, "User /\nFrontend", "Chat API\n(Express)", "POST /api/chat {message}"); y+=step
msg(y, "Chat API\n(Express)", "Auth +\nRate Limit", "verify JWT + rate-limit"); y+=step
msg(y, "Auth +\nRate Limit", "Chat API\n(Express)", "user / role OK", color=TEAL_D, ret=True); y+=step
msg(y, "Chat API\n(Express)", "MongoDB", "ensure session + save user msg"); y+=step
msg(y, "Chat API\n(Express)", "Cache\n(Redis)", "lookup cached answer"); y+=step
msg(y, "Cache\n(Redis)", "Chat API\n(Express)", "miss", color=AMBER, ret=True); y+=step
msg(y, "Chat API\n(Express)", "Queue\n(Bull)", "enqueue chat job"); y+=step
msg(y, "Queue\n(Bull)", "n8n AI\nAgent", "POST role-scoped payload"); y+=step
msg(y, "n8n AI\nAgent", "MongoDB", "role-scoped query"); y+=step
msg(y, "MongoDB", "n8n AI\nAgent", "rows / aggregates", color=TEAL_D, ret=True); y+=step
msg(y, "n8n AI\nAgent", "Chat API\n(Express)", "reply + data + files", color=TEAL_D, ret=True); y+=step
msg(y, "Chat API\n(Express)", "MongoDB", "persist agent msg + cache set"); y+=step
msg(y, "Chat API\n(Express)", "User /\nFrontend", "200 {reply, data, sources}", color=TEAL_D, ret=True)

# ============================================================================
# 14. DATABASE DESIGN
# ============================================================================
s = new_slide(WHITE)
titlebar(s, "Database Design", "MongoDB — Mongoose")
textbox(s, Inches(0.6), Inches(1.25), Inches(12.2), Inches(0.55),
        "MongoDB document store with ~60 Mongoose collections. Below: the core entities and the Chat-module schema "
        "(referenced relationships shown).", 12.5, DARK)
def entity(x, y, title, fields, w=Inches(2.75), header=NAVY):
    hh=Inches(0.42)
    box(s, x, y, w, hh, title, header, WHITE, 12, bold=True, shape=MSO_SHAPE.RECTANGLE)
    fh=Inches(0.27)*len(fields)
    body=box(s, x, y+hh, w, fh, "", WHITE, DARK, 10, shape=MSO_SHAPE.RECTANGLE, line=NAVY, line_w=0.75)
    tf=body.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_top=Inches(0.04); tf.margin_left=Inches(0.08)
    for i,f in enumerate(fields):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        r=p.add_run(); r.text=f
        r.font.size=Pt(9.5); r.font.color.rgb=DARK; r.font.name="Consolas"
        if f.startswith("PK") or f.startswith("FK"):
            r.font.color.rgb=TEAL_D; r.font.bold=True
    return (x,y,w,hh+fh)
# Core ring (top)
entity(Inches(0.55), Inches(1.95), "User", ["PK _id","name, email","passwordHash","FK role","status"], header=TEAL_D)
entity(Inches(3.55), Inches(1.95), "Role / Permission", ["PK _id","name","permissions[]","scopes[]"])
entity(Inches(6.55), Inches(1.95), "Student / Parent", ["PK _id","FK user","FK class","guardianInfo"])
entity(Inches(9.55), Inches(1.95), "Domain Collections", ["Attendance","Exam / ReportCard","Fee / Transport","Timetable, Homework","Complaint, Notification"], header=AMBER)
# Chat module (bottom)
e_sess=entity(Inches(0.55), Inches(4.35), "ChatSession", ["PK _id","FK user","sessionKey","roleAtCreation","allowedScopes[]","status, lastMessageAt"], header=NAVY)
e_msg=entity(Inches(4.0), Inches(4.35), "ChatMessage", ["PK _id","FK session","senderType","FK senderUser","text, data","actions[], sources[]"], header=NAVY)
e_file=entity(Inches(7.45), Inches(4.35), "ChatFile", ["PK _id","FK session","FK owner","filename, mimeType","data (Buffer)","expiresAt (TTL)"], header=NAVY)
e_md=entity(Inches(10.5), Inches(4.35), "Message / Delivery", ["PK _id","FK sender, recipient","text, attachments","status (sent/","  delivered/read)"], header=TEAL_D)
# relationships
connector(s, Inches(1.9), Inches(2.4), Inches(3.55), Inches(2.4), color=GREY, arrow=True); edge_label(s, Inches(2.0), Inches(2.1), "has role", GREY, 8, Inches(1.2))
connector(s, Inches(1.5), Inches(3.5), Inches(1.5), Inches(4.35), color=GREY, arrow=True); edge_label(s, Inches(1.55), Inches(3.85), "owns sessions", GREY, 8, Inches(1.5))
connector(s, Inches(3.3), Inches(5.2), Inches(4.0), Inches(5.2), color=GREY, arrow=True); edge_label(s, Inches(3.0), Inches(5.45), "1 → N msgs", GREY, 8, Inches(1.3))
connector(s, Inches(6.75), Inches(5.2), Inches(7.45), Inches(5.2), color=GREY, arrow=True); edge_label(s, Inches(6.55), Inches(5.45), "attaches", GREY, 8, Inches(1.2))
textbox(s, Inches(0.6), Inches(6.85), Inches(12.2), Inches(0.5),
        "Keys: PK = primary key (_id), FK = referenced ObjectId. Indexes: ChatSession(user, sessionKey) unique, "
        "ChatMessage(session, createdAt), ChatFile TTL on expiresAt, MessageDelivery(message, recipient) unique.",
        10.5, GREY)

# ============================================================================
# 15. VALIDATION & TESTING
# ============================================================================
s = new_slide(LIGHT)
titlebar(s, "Validation & Testing", "QUALITY ASSURANCE")
textbox(s, Inches(0.7), Inches(1.4), Inches(12), Inches(0.5),
        "A multi-layer strategy validates that EduNexus meets its functional and non-functional requirements.", 14, DARK)
cols = [
    ("Validation", TEAL_D, [
        "Input validation (empty / length / type) on every API.",
        "Schema validation via Mongoose models.",
        "JWT + role checks validate access on each request.",
        "Requirement traceability — each FR mapped to a test.",
    ]),
    ("Testing Levels", NAVY, [
        "Unit testing — controllers, services, utils.",
        "Integration testing — API + DB + queue.",
        "System testing — end-to-end role workflows.",
        "User Acceptance Testing with sample school data.",
    ]),
    ("Techniques & Tools", AMBER, [
        "Black-box (functional) + White-box (unit).",
        "Manual exploratory + Postman API tests.",
        "Performance: 100 concurrent users, <90s AI replies.",
        "Security: auth, RBAC & rate-limit verification.",
    ]),
]
cx=Inches(0.7); cw=Inches(4.0)
for title_c, color_c, items in cols:
    box(s, cx, Inches(2.1), cw, Inches(0.55), title_c, color_c, WHITE, 15, bold=True)
    bx=box(s, cx, Inches(2.7), cw, Inches(3.6), "", WHITE, DARK, 12, line=RGBColor(0xD5,0xDB,0xE5))
    tf=bx.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.12); tf.margin_top=Inches(0.1)
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_after=Pt(8)
        r=p.add_run(); r.text="✓ "; r.font.bold=True; r.font.color.rgb=color_c; r.font.size=Pt(12.5)
        r2=p.add_run(); r2.text=it; r2.font.size=Pt(12.5); r2.font.color.rgb=DARK; r2.font.name="Segoe UI"
    cx+=cw+Inches(0.25)
textbox(s, Inches(0.7), Inches(6.5), Inches(12), Inches(0.6),
        "Outcome: core modules pass functional tests; AI chat verified for correct role-scoping, caching, queue fallback and error handling.",
        13, NAVY, bold=True)

# ============================================================================
# 16. TEST CASES (Black-box + Unit)
# ============================================================================
def add_table(slide, rows, x, y, w, col_w, header_fill=NAVY, size=10):
    nrows=len(rows); ncols=len(rows[0])
    gt=slide.shapes.add_table(nrows, ncols, x, y, w, Inches(0.3)*nrows).table
    for j,cw in enumerate(col_w):
        gt.columns[j].width=cw
    for i,row in enumerate(rows):
        for j,val in enumerate(row):
            cell=gt.cell(i,j)
            cell.margin_left=Inches(0.05); cell.margin_right=Inches(0.05)
            cell.margin_top=Inches(0.02); cell.margin_bottom=Inches(0.02)
            tf=cell.text_frame; tf.word_wrap=True
            p=tf.paragraphs[0]; r=p.add_run(); r.text=val
            r.font.size=Pt(size); r.font.name="Segoe UI"
            if i==0:
                cell.fill.solid(); cell.fill.fore_color.rgb=header_fill
                r.font.color.rgb=WHITE; r.font.bold=True
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if i%2 else BOXBLUE
                r.font.color.rgb=DARK
    return gt

# Black-box
s = new_slide(LIGHT)
titlebar(s, "Black-Box Test Cases — Chat Module", "FUNCTIONAL TESTING")
bb = [
    ["ID","Test Case","Input","Expected Result","Status"],
    ["TC-01","Send valid query","Auth user posts \"Show today's attendance\"","200 OK; role-scoped reply with data returned","Pass"],
    ["TC-02","Empty message","message = \"\"","400 Bad Request — \"Message cannot be empty\"","Pass"],
    ["TC-03","Oversized message",">2000 characters","400 Bad Request — length error","Pass"],
    ["TC-04","Unauthenticated request","No / invalid JWT","401 Unauthorized; access denied","Pass"],
    ["TC-05","Rate limit","3rd request within 30s","429 Too Many Requests","Pass"],
    ["TC-06","Role scoping","Parent asks for staff salaries","Agent returns no out-of-scope data","Pass"],
    ["TC-07","Cached query","Repeat same query in TTL","Fast cached response served","Pass"],
    ["TC-08","Resume session","Open existing sessionKey","Prior messages loaded in order","Pass"],
    ["TC-09","Download attachment","GET valid file owned by user","File streamed with correct mime-type","Pass"],
    ["TC-10","Agent failure","n8n unreachable / 5xx","Graceful error message; no crash","Pass"],
]
add_table(s, bb, Inches(0.45), Inches(1.4), Inches(12.4),
          [Inches(0.9),Inches(2.4),Inches(3.0),Inches(4.6),Inches(0.9)], size=10.5)

# Unit
s = new_slide(LIGHT)
titlebar(s, "Unit Test Cases — Chat Module", "WHITE-BOX TESTING")
ut = [
    ["ID","Unit / Function","Scenario","Expected Output"],
    ["UT-01","ensureSession()","New user, no sessionId","Creates ChatSession with role & sessionKey = userId"],
    ["UT-02","ensureSession()","Existing sessionKey","Returns same session; updates lastMessageAt"],
    ["UT-03","sendMessage() validation","message is null/blank","Returns 400, no DB write"],
    ["UT-04","persistAttachments()","base64 within size limit","Creates ChatFile; returns normalized url"],
    ["UT-05","persistAttachments()","Attachment > max bytes","Throws \"exceeds maximum allowed size\""],
    ["UT-06","processChat() extractReply","n8n returns [{output:'x'}]","Extracts reply = 'x'"],
    ["UT-07","processChat() extractReply","Empty / null body","Returns agentError reply, no throw"],
    ["UT-08","processChat() error path","Axios 4xx response","Returns agentDenied with status code"],
    ["UT-09","getRoleAccessProfile()","role = 'parent'","Returns parent-scoped access profile"],
    ["UT-10","cache get/set","set then get within TTL","Returns stored response object"],
    ["UT-11","downloadAttachment()","File not owned by user","Returns 404 (ownership enforced)"],
]
add_table(s, ut, Inches(0.55), Inches(1.45), Inches(12.2),
          [Inches(0.95),Inches(2.9),Inches(3.2),Inches(5.15)], size=11, header_fill=TEAL_D)

# ============================================================================
# CLOSING
# ============================================================================
s = new_slide(NAVY)
textbox(s, Inches(1), Inches(2.9), Inches(11.3), Inches(1.0), "Thank You", 50, WHITE, bold=True, align=PP_ALIGN.CENTER)
textbox(s, Inches(1), Inches(4.0), Inches(11.3), Inches(0.7),
        "EduNexus — An AI-Powered School Management System", 20, TEAL, bold=True, align=PP_ALIGN.CENTER)
textbox(s, Inches(1), Inches(4.8), Inches(11.3), Inches(0.6),
        "Questions & Demo", 16, RGBColor(0xB8,0xC2,0xD4), align=PP_ALIGN.CENTER)

out = r"d:\fyp\edunexus\EduNexus_Project_Evaluation.pptx"
try:
    prs.save(out)
except PermissionError:
    out = r"d:\fyp\edunexus\EduNexus_Project_Evaluation_new.pptx"
    prs.save(out)
print("SAVED:", out, "slides:", len(prs.slides._sldIdLst))
